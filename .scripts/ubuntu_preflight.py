#!/usr/bin/env python3
"""Read-only deployment preflight for the supported Ubuntu runtime."""
from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Callable, Mapping

import platform_compat


REPO = Path(__file__).resolve().parent.parent
MIN_PYTHON = (3, 10)
PYTHON_IMPORTS = {
    "PyYAML": "yaml",
    "numpy": "numpy",
    "scipy": "scipy",
    "matplotlib": "matplotlib",
    "networkx": "networkx",
    "rank-bm25": "rank_bm25",
    "tiktoken": "tiktoken",
    "requests": "requests",
    "python-dotenv": "dotenv",
    "PyMuPDF": "fitz",
    "Pillow": "PIL",
    "python-pptx": "pptx",
    "openpyxl": "openpyxl",
    "xlrd": "xlrd",
    "xlwt": "xlwt",
}
SYSTEM_COMMANDS = {
    "git": ("git",),
    "curl": ("curl",),
    "ripgrep": ("rg",),
    "tesseract": ("tesseract",),
}


def _check(name: str, ok: bool, detail: str) -> dict[str, object]:
    return {"name": name, "ok": ok, "detail": detail}


def _probe_import(module: str) -> tuple[bool, str]:
    with tempfile.TemporaryDirectory(prefix="asks-python-import-") as temp_dir:
        env = os.environ.copy()
        env.update({
            "PYTHONDONTWRITEBYTECODE": "1",
            "MPLCONFIGDIR": temp_dir,
            "XDG_CACHE_HOME": temp_dir,
            "XDG_CONFIG_HOME": temp_dir,
        })
        try:
            result = subprocess.run(
                [
                    sys.executable,
                    "-I",
                    "-B",
                    "-c",
                    f"import importlib; importlib.import_module({module!r})",
                ],
                capture_output=True,
                text=True,
                timeout=60,
                check=False,
                env=env,
            )
        except (OSError, subprocess.SubprocessError) as exc:
            return False, str(exc)
    if result.returncode == 0:
        return True, module
    detail = result.stderr.strip().splitlines()
    return False, detail[-1] if detail else f"import {module} failed"


def _python_checks(
    *,
    version_info=sys.version_info,
    import_probe: Callable[[str], tuple[bool, str]] = _probe_import,
) -> list[dict[str, object]]:
    version = tuple(version_info[:3])
    checks = [
        _check(
            "python-version",
            version >= MIN_PYTHON,
            ".".join(str(value) for value in version),
        )
    ]
    for distribution, module in PYTHON_IMPORTS.items():
        available, detail = import_probe(module)
        checks.append(_check(
            f"python:{distribution}",
            available,
            detail,
        ))
    return checks


def _tesseract_language_check(
    binary: Path,
    *,
    run: Callable[..., subprocess.CompletedProcess] = subprocess.run,
) -> dict[str, object]:
    try:
        result = run(
            [str(binary), "--list-langs"],
            capture_output=True,
            text=True,
            timeout=30,
            check=False,
        )
    except (OSError, subprocess.SubprocessError) as exc:
        return _check("tesseract-languages", False, str(exc))
    languages = {
        line.strip() for line in result.stdout.splitlines()
        if line.strip() and not line.lower().startswith("list of available")
    }
    required = {"eng", "chi_sim"}
    missing = sorted(required - languages)
    return _check(
        "tesseract-languages",
        result.returncode == 0 and not missing,
        ", ".join(sorted(languages)) if not missing else f"missing {', '.join(missing)}",
    )


def _system_checks(
    *,
    which: platform_compat.Which = shutil.which,
    environ: Mapping[str, str] | None = None,
    platform: str | None = None,
    run: Callable[..., subprocess.CompletedProcess] = subprocess.run,
) -> tuple[list[dict[str, object]], Path | None]:
    checks: list[dict[str, object]] = []
    resolved: dict[str, Path | None] = {}
    for label, names in SYSTEM_COMMANDS.items():
        executable = platform_compat.find_executable(*names, which=which)
        resolved[label] = executable
        checks.append(_check(
            f"command:{label}", executable is not None,
            str(executable) if executable else f"missing {'/'.join(names)}",
        ))

    soffice = platform_compat.find_soffice(environ=environ, which=which)
    checks.append(_check(
        "command:libreoffice",
        soffice is not None,
        str(soffice) if soffice else "missing soffice/libreoffice",
    ))

    openers = platform_compat.default_open_commands(
        ".", platform=platform, which=which
    )
    checks.append(_check(
        "desktop-opener",
        bool(openers),
        ", ".join(command[0] for command in openers) if openers else "missing opener",
    ))
    trash_commands = platform_compat.system_trash_commands(
        ".", platform=platform, which=which
    )
    checks.append(_check(
        "system-trash",
        bool(trash_commands),
        ", ".join(command[0] for command in trash_commands)
        if trash_commands else "project temp/trash fallback only",
    ))
    if resolved["tesseract"] is not None:
        checks.append(_tesseract_language_check(resolved["tesseract"], run=run))
    else:
        checks.append(_check(
            "tesseract-languages", False, "tesseract is unavailable"
        ))
    return checks, soffice


def _docling_check(
    *,
    repo: Path = REPO,
    run: Callable[..., subprocess.CompletedProcess] = subprocess.run,
) -> dict[str, object]:
    binary = repo / ".venv-docling/bin/python"
    if not binary.is_file() or not os.access(binary, os.X_OK):
        return _check("docling-environment", False, f"missing {binary}")
    try:
        result = run(
            [str(binary), "-c", "import docling"],
            capture_output=True,
            text=True,
            timeout=60,
            check=False,
        )
    except (OSError, subprocess.SubprocessError) as exc:
        return _check("docling-environment", False, str(exc))
    detail = str(binary) if result.returncode == 0 else (
        result.stderr.strip() or "docling import failed"
    )
    return _check("docling-environment", result.returncode == 0, detail)


def _libreoffice_smoke(soffice: Path) -> dict[str, object]:
    sentinel = "Ran-ASKS Ubuntu smoke"
    try:
        import fitz
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        return _check("libreoffice-smoke", False, f"missing Python dependency: {exc}")

    try:
        with tempfile.TemporaryDirectory(prefix="asks-ubuntu-smoke-") as temp_dir:
            workspace = Path(temp_dir)
            source = workspace / "smoke.pptx"
            output_dir = workspace / "output"
            profile_dir = workspace / "profile"
            output_dir.mkdir()
            profile_dir.mkdir()

            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            text_box = slide.shapes.add_textbox(
                Inches(1), Inches(1), Inches(8), Inches(1)
            )
            text_box.text_frame.text = sentinel
            presentation.save(source)

            result = subprocess.run(
                [
                    str(soffice),
                    f"-env:UserInstallation={profile_dir.as_uri()}",
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(output_dir),
                    str(source),
                ],
                capture_output=True,
                text=True,
                timeout=180,
                check=False,
                env=platform_compat.soffice_environment(soffice),
            )
            rendered = output_dir / "smoke.pdf"
            if result.returncode != 0 or not rendered.is_file():
                detail = result.stderr.strip() or result.stdout.strip() or "PDF not created"
                return _check("libreoffice-smoke", False, detail)
            with fitz.open(rendered) as document:
                text = "\n".join(page.get_text() for page in document)
                ok = len(document) == 1 and sentinel in text
            return _check(
                "libreoffice-smoke", ok,
                "PPTX -> PDF and text verified" if ok else "rendered PDF lost sentinel text",
            )
    except (OSError, subprocess.SubprocessError, ValueError) as exc:
        return _check("libreoffice-smoke", False, str(exc))


def build_report(
    *,
    python_only: bool = False,
    include_docling: bool = False,
    smoke: bool = False,
    which: platform_compat.Which = shutil.which,
    environ: Mapping[str, str] | None = None,
    platform: str | None = None,
    run: Callable[..., subprocess.CompletedProcess] = subprocess.run,
    import_probe: Callable[[str], tuple[bool, str]] = _probe_import,
    version_info=sys.version_info,
    repo: Path = REPO,
) -> dict[str, object]:
    checks = _python_checks(version_info=version_info, import_probe=import_probe)
    soffice: Path | None = None
    if not python_only:
        system_checks, soffice = _system_checks(
            which=which, environ=environ, platform=platform, run=run
        )
        checks.extend(system_checks)
    if include_docling:
        checks.append(_docling_check(repo=repo, run=run))
    if smoke:
        if python_only:
            checks.append(_check(
                "libreoffice-smoke", False,
                "--smoke cannot be combined with --python-only",
            ))
        elif soffice is None:
            checks.append(_check(
                "libreoffice-smoke", False, "LibreOffice is unavailable"
            ))
        else:
            checks.append(_libreoffice_smoke(soffice))
    return {
        "schema": "ubuntu-preflight-v1",
        "ready": all(bool(item["ok"]) for item in checks),
        "checks": checks,
    }


def _print_report(report: dict[str, object]) -> None:
    print(f"Ubuntu preflight: {'READY' if report['ready'] else 'INCOMPLETE'}")
    for item in report["checks"]:
        marker = "ready" if item["ok"] else "missing"
        print(f"[{marker}] {item['name']}: {item['detail']}")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Read-only Ubuntu runtime and deployment preflight"
    )
    parser.add_argument(
        "--python-only", action="store_true",
        help="check the Python version and requirements imports only",
    )
    parser.add_argument(
        "--docling", action="store_true",
        help="also verify the isolated .venv-docling environment",
    )
    parser.add_argument(
        "--smoke", action="store_true",
        help="render a temporary one-slide PPTX through LibreOffice and verify its PDF",
    )
    parser.add_argument(
        "--strict", action="store_true",
        help="exit nonzero when any requested check is unavailable or fails",
    )
    parser.add_argument("--json", action="store_true", help="emit JSON")
    args = parser.parse_args()

    report = build_report(
        python_only=args.python_only,
        include_docling=args.docling,
        smoke=args.smoke,
    )
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        _print_report(report)
    if args.strict and not report["ready"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
